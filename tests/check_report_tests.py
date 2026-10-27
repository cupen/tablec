#!/usr/bin/env python3
"""Integration tests for `tablec check --report` (openspec change add-validation-report).

Drives the compiled CLI binary (target/debug/tablec) against generated Excel
fixtures and asserts the validation-report behavior from the spec deltas:

- report format selected by extension (.json / .html), others rejected
- reports written on every exit path (clean, errors, no-files) and printed
- JSON envelope members, summary counters, diagnostics in pipeline order
- HTML self-containment, seeded diagnostics, severity badges, passed sheets
- no report file without --report
- .xml extension fails fast before scanning
- unwritable report path fails the command naming the path

Run:  python3 tests/check_report_tests.py
Needs openpyxl for fixture generation:  python3 -m pip install --user openpyxl
"""

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("error: openpyxl is required to generate xlsx fixtures "
          "(python3 -m pip install --user openpyxl)")
    sys.exit(2)

REPO_ROOT = Path(__file__).resolve().parent.parent
BIN = REPO_ROOT / "target" / "debug" / "tablec"

RED = '\033[91m'
GREEN = '\033[92m'
YELLOW = '\033[93m'
BLUE = '\033[94m'
RESET = '\033[0m'
BOLD = '\033[1m'

ENVELOPE_KEYS = {
    "schema_version", "tablec_version", "generated_at",
    "input_dir", "parser", "summary", "diagnostics",
}
SUMMARY_KEYS = {"errors", "warnings", "files_checked", "tables_checked"}
RFC3339_UTC = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$")


# ---------------------------------------------------------------------------
# Fixture generation (standard 5-row tablec layout: names, types, comments,
# field constraints, table constraints, data from row 6)
# ---------------------------------------------------------------------------

def add_sheet(wb, name, columns, rows):
    """columns: [(field, type, constraint)], rows: list of value-lists."""
    ws = wb.create_sheet(name)
    for j, (field, ty, constraint) in enumerate(columns, 1):
        ws.cell(row=1, column=j, value=field)
        ws.cell(row=2, column=j, value=ty)
        ws.cell(row=3, column=j, value="")
        ws.cell(row=4, column=j, value=constraint)
        ws.cell(row=5, column=j, value="")
    for i, row in enumerate(rows, 6):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=v)


def make_seeded_dir(root):
    """Table set with two seeded errors plus one clean sheet.

    - bad_range.xlsx/Range: int8 value 200 (out of range) -> ValueOutOfRange
    - drops.xlsx/Drop: @ref("Item.id") with a missing target -> FK violation
    - items.xlsx/Item: clean
    """
    root.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Range", [("hp", "int8", "")], [[200]])
    wb.save(root / "bad_range.xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Item", [("id", "int8", "")], [[1], [2]])
    wb.save(root / "items.xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Drop", [("item_id", "int8", '@ref("Item.id")')], [[1], [99]])
    wb.save(root / "drops.xlsx")
    return root


def make_clean_dir(root):
    """Fully clean table set: two files, two sheets."""
    root.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Alpha", [("id", "int8", "")], [[1], [2], [3]])
    wb.save(root / "a_alpha.xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Beta", [("id", "int8", ""), ("name", "string", "")],
              [[10, "x"], [20, "y"]])
    wb.save(root / "b_beta.xlsx")
    return root


def ensure_binary():
    if BIN.exists():
        return BIN
    env = dict(os.environ)
    env["PATH"] = os.path.join(Path.home(), ".cargo", "bin") + os.pathsep + env.get("PATH", "")
    print(f"{YELLOW}binary not found, running cargo build ...{RESET}")
    subprocess.run(["cargo", "build"], cwd=REPO_ROOT, env=env, check=True)
    if not BIN.exists():
        print(f"error: {BIN} still missing after cargo build")
        sys.exit(2)
    return BIN


def run_check(args, timeout=120):
    return subprocess.run([str(BIN), "check"] + args,
                          capture_output=True, text=True, timeout=timeout)


def external_ref_violations(html):
    """Occurrences of src=/href=/url( pointing at absolute remote resources."""
    bad = []
    for marker in ("src=", "href=", "url("):
        for m in re.finditer(re.escape(marker), html):
            tail = html[m.end():].lstrip()
            tail = re.sub(r'^["\']', "", tail)
            if tail.startswith("http://") or tail.startswith("https://") or tail.startswith("//"):
                bad.append(marker + tail[:30])
    return bad


class TestResult:
    def __init__(self, name, passed, message=""):
        self.name = name
        self.passed = passed
        self.message = message


FAILURES = []


def check(cond, label):
    """Record one assertion; returns True when it held."""
    if not cond:
        FAILURES.append(label)
    return cond


# ---------------------------------------------------------------------------
# Test cases
# ---------------------------------------------------------------------------

def test_seeded_error_run_reports(tmp):
    """Error run with both report formats: exit non-zero, JSON envelope and
    counters correct, HTML shows the seeded diagnostics, badges and the
    passed sheet; both written paths printed."""
    seeded = make_seeded_dir(tmp / "seeded")
    out = tmp / "seeded_out"
    json_path = out / "report.json"
    html_path = out / "report.html"

    r = run_check([str(seeded), "--report", str(json_path), "--report", str(html_path)])
    check(r.returncode != 0, "seeded-error run must exit non-zero")
    check(f"Report written: {json_path}" in r.stdout, f"stdout must name {json_path}")
    check(f"Report written: {html_path}" in r.stdout, f"stdout must name {html_path}")

    # --- JSON report ---
    if check(json_path.exists(), "JSON report must exist"):
        data = json.loads(json_path.read_text())
        check(set(data) == ENVELOPE_KEYS,
              f"envelope must have exactly the v1 members, got {sorted(data)}")
        check(data["schema_version"] == 1, "schema_version must be 1")
        check(isinstance(data["tablec_version"], str) and data["tablec_version"],
              "tablec_version must be a non-empty string")
        check(bool(RFC3339_UTC.match(data["generated_at"])),
              f"generated_at must be RFC 3339 UTC, got {data['generated_at']!r}")
        check(data["input_dir"] == str(seeded), "input_dir must name the scanned directory")
        check(data["parser"] == "standard", "parser must be the standard parser")
        summary = data["summary"]
        check(set(summary) == SUMMARY_KEYS,
              f"summary must have exactly the v1 members, got {sorted(summary)}")
        check(summary["errors"] == 2, f"errors must be 2, got {summary['errors']}")
        check(summary["warnings"] == 0, f"warnings must be 0, got {summary['warnings']}")
        check(summary["files_checked"] == 3,
              f"files_checked must count all 3 files incl. the failing one, got {summary['files_checked']}")
        check(summary["tables_checked"] == 2,
              f"tables_checked must count the 2 parsed sheets, got {summary['tables_checked']}")

        diags = data["diagnostics"]
        check(len(diags) == 2, f"both seeded errors must be reported, got {len(diags)}")
        if len(diags) == 2:
            check([d["code"] for d in diags] == ["ValueOutOfRange", "ConstraintForeignKeyViolation"],
                  "diagnostics must appear in pipeline order (per-table, then cross-table)")
            check(all(d["severity"] == "Error" for d in diags),
                  "seeded diagnostics are errors")
            check(diags[0]["message"] == "value '200' not in int8 range [-128, 127]",
                  f"unexpected first message: {diags[0]['message']!r}")
            check(diags[0]["location"]["file"] == str(seeded / "bad_range.xlsx")
                  and diags[0]["location"]["sheet"] == "Range",
                  f"first diagnostic must locate the seeded sheet, got {diags[0]['location']}")
            for d in diags:
                check(set(d) == {"severity", "code", "message", "location"},
                      "diagnostic entries must reuse the core serialization shape")
                check(set(d["location"]) == {"file", "sheet", "line", "column"},
                      "location must reuse the core SourceLocation shape")

    # --- HTML report ---
    if check(html_path.exists(), "HTML report must exist"):
        html = html_path.read_text()
        check(html.lstrip().startswith("<!DOCTYPE html>"), "HTML must be a complete document")
        check(external_ref_violations(html) == [],
              "HTML must be self-contained: no external src/href/url references")
        for label in ["校验报告", "问题明细", "通过的表", "错误", "警告", "生成时间", "工具版本", "解析器", "输入目录"]:
            check(label in html, f"HTML must carry Chinese UI label {label!r}")
        check("value &#39;200&#39; not in int8 range [-128, 127]" in html,
              "HTML must contain the seeded int8 diagnostic message")
        check("@ref(Item.id): row 2: host &#39;item_id&#39; = 99 missing from target Item.id" in html,
              "HTML must contain the seeded @ref diagnostic message")
        check('class="badge badge-error"' in html, "HTML must carry error severity badges")
        check('data-severity="error"' in html, "diagnostic entries must carry a severity data attribute")
        check('data-filter="error"' in html and 'data-filter="warning"' in html,
              "HTML must offer the severity filter")
        check("bad_range.xlsx" in html, "diagnostics must be grouped under the file name")
        check("Range" in html, "diagnostics must be grouped under the sheet name")
        check("<li>Item</li>" in html, "the clean sheet must be listed as passed")
        check("<li>Drop</li>" not in html,
              "the sheet carrying the @ref violation must NOT be listed as passed")
        # Both artifacts describe the same run: HTML counters match the JSON.
        m = re.search(r'class="stat stat-errors"><span class="num">(\d+)</span>', html)
        check(m is not None and m.group(1) == "2", "HTML error stat must match summary.errors")
        m = re.search(r'class="stat stat-warnings"><span class="num">(\d+)</span>', html)
        check(m is not None and m.group(1) == "0", "HTML warning stat must match summary.warnings")


def test_clean_run_report(tmp):
    """Clean run: exit zero, zero counts, files_checked > 0, HTML empty state
    and passed sheets."""
    clean = make_clean_dir(tmp / "clean")
    out = tmp / "clean_out"
    json_path = out / "r.json"
    html_path = out / "r.html"

    r = run_check([str(clean), "--report", str(json_path), "--report", str(html_path)])
    check(r.returncode == 0, f"clean run must exit zero, got {r.returncode}: {r.stderr}")
    check(json_path.exists() and html_path.exists(), "both reports must exist on a clean run")

    if json_path.exists():
        data = json.loads(json_path.read_text())
        s = data["summary"]
        check(s["errors"] == 0 and s["warnings"] == 0, "clean run must record zero errors and warnings")
        check(s["files_checked"] > 0, f"clean run must record files_checked > 0, got {s['files_checked']}")
        check(s["files_checked"] == 2 and s["tables_checked"] == 2,
              f"clean fixture has 2 files / 2 sheets, got {s}")
        check(data["diagnostics"] == [], "clean run must have an empty diagnostics array")

    if html_path.exists():
        html = html_path.read_text()
        check("未发现问题，全部通过。" in html, "clean run must show the distinct empty state")
        check("<li>Alpha</li>" in html and "<li>Beta</li>" in html,
              "clean run must list all sheets as passed")


def test_no_flag_writes_nothing(tmp):
    """Without --report no report file is created anywhere."""
    clean = make_clean_dir(tmp / "noflag")
    r = run_check([str(clean)])
    check(r.returncode == 0, "clean run without --report must still exit zero")
    strays = [str(p.relative_to(tmp)) for p in tmp.rglob("*")
              if p.suffix in (".json", ".html") and p.is_file()]
    check(strays == [], f"no report file may be created without --report, found {strays}")


def test_xml_extension_fails_fast(tmp):
    """A .xml report request exits non-zero naming the supported extensions,
    before any file scan and without writing a report."""
    seeded = make_seeded_dir(tmp / "seeded2")
    out = tmp / "xml_out"
    bad = out / "report.xml"

    r = run_check([str(seeded), "--report", str(bad)])
    check(r.returncode != 0, ".xml report request must exit non-zero")
    check(".json" in r.stderr and ".html" in r.stderr,
          f"error must name the supported extensions, got: {r.stderr!r}")
    check(str(bad) in r.stderr, f"error must name the offending path, got: {r.stderr!r}")
    check(not bad.exists(), "no report file may be written for a rejected extension")
    # Pre-validation: the run must fail before scanning (no per-file findings,
    # no completion banner, no seeded diagnostics rendered).
    check("ValueOutOfRange" not in r.stderr, "no table file may be scanned before validation fails")
    check("Check finished successfully" not in r.stdout and "Found" not in r.stdout,
          "no scan output may appear before validation fails")

    # Extension validation also precedes target-existence handling.
    r2 = run_check([str(tmp / "does_not_exist"), "--report", str(out / "x.xml")])
    check(r2.returncode != 0, ".xml request on a missing target must still fail on the extension")
    check("No Excel files found" not in r2.stdout,
          "extension validation must precede the target-existence notice")


def test_unwritable_report_path_fails(tmp):
    """A report path under a regular file fails the command with the path
    named, even though the check itself is clean."""
    clean = make_clean_dir(tmp / "clean2")
    blocker = tmp / "blocker.json"
    blocker.write_text("i am a regular file")
    report_path = blocker / "report.json"

    r = run_check([str(clean), "--report", str(report_path)])
    check(r.returncode != 0,
          f"unwritable report path must fail the command even on a clean check, got {r.returncode}")
    output = r.stdout + r.stderr
    check("failed to write report" in output, f"error must identify the write failure: {output!r}")
    check(str(report_path) in output, f"error must name the unwritable path: {output!r}")


def test_report_overwrites_existing_file(tmp):
    """An existing file at the report path is replaced by the new run."""
    clean = make_clean_dir(tmp / "clean3")
    json_path = tmp / "clean3" / "stale.json"
    json_path.write_text("stale junk")

    r = run_check([str(clean), "--report", str(json_path)])
    check(r.returncode == 0, f"clean run must succeed, got {r.returncode}: {r.stderr}")
    if json_path.exists():
        data = json.loads(json_path.read_text())
        check(data.get("schema_version") == 1, "stale file must be replaced by the new report")


def test_no_files_run_report(tmp):
    """A directory without matching spreadsheet files still writes a report
    reflecting the run (one warning, zero files/tables)."""
    empty = tmp / "empty_dir"
    empty.mkdir()
    json_path = tmp / "empty_out" / "r.json"

    r = run_check([str(empty), "--report", str(json_path)])
    check(r.returncode == 0, "no-matching-files run is not an error")
    if json_path.exists():
        data = json.loads(json_path.read_text())
        s = data["summary"]
        check(s["warnings"] == 1 and s["errors"] == 0,
              f"no-files run must carry the no-files warning, got {s}")
        check(s["files_checked"] == 0 and s["tables_checked"] == 0,
              f"no-files run must record zero files/tables, got {s}")


TESTS = [
    test_seeded_error_run_reports,
    test_clean_run_report,
    test_no_flag_writes_nothing,
    test_xml_extension_fails_fast,
    test_unwritable_report_path_fails,
    test_report_overwrites_existing_file,
    test_no_files_run_report,
]


def main():
    print(f"\n{BOLD}{BLUE}=== tablec check --report integration tests ==={RESET}\n")
    ensure_binary()

    results = []
    for test in TESTS:
        with tempfile.TemporaryDirectory(prefix="tablec_report_") as td:
            try:
                test(Path(td))
                passed, msg = True, ""
            except Exception as e:  # noqa: BLE001 - report any failure as a test failure
                passed, msg = False, f"{type(e).__name__}: {e}"
        results.append(TestResult(test.__name__, passed, msg))
        status = f"{GREEN}PASS{RESET}" if passed else f"{RED}FAIL{RESET}"
        print(f"  [{status}] {test.__name__}")
        if not passed:
            for f in FAILURES:
                print(f"      {RED}- {f}{RESET}")
        FAILURES.clear()

    failed = sum(1 for r in results if not r.passed)
    print(f"\n{BOLD}Total: {len(results)} tests  "
          f"{GREEN}passed: {len(results) - failed}{RESET}  "
          f"{RED if failed else GREEN}failed: {failed}{RESET}")
    if shutil.which("cargo") is None:
        print(f"{YELLOW}note: cargo not on PATH; binary reuse only "
              f"(set PATH=\"$HOME/.cargo/bin:$PATH\" when building){RESET}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
